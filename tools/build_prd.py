#!/usr/bin/env python3
"""Build the persistent Open Bikeshare MCP PRD from its Markdown source.

The builder starts from the retained original DOCX so the revised document
inherits its page system, styles, theme, numbering, header, and footer.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = PROJECT_ROOT / "docs" / "Open_Bikeshare_MCP_PRD.md"
DEFAULT_TEMPLATE = (
    PROJECT_ROOT
    / "docs"
    / "reference"
    / "Open_Bikeshare_MCP_PRD_original.docx"
)
DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "docs"
    / "Open_Bikeshare_MCP_PRD_Cloudflare_v1.3.docx"
)

NAVY = "14233B"
BLUE = "2357C6"
SLATE = "5E6B7A"
PALE_BLUE = "E5EDF9"
ROW_ALT = "F3F5F7"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360


TABLE_WIDTHS_BY_HEADERS: dict[tuple[str, ...], list[int]] = {
    ("Field", "Value"): [1800, 7560],
    ("Decision", "V1 choice", "Rationale"): [1800, 2700, 4860],
    ("Persona", "Need", "Primary job"): [1800, 3000, 4560],
    ("Metric", "Beta target", "Measurement"): [2100, 3000, 4260],
    ("Tool", "Purpose", "Typical use"): [2400, 4200, 2760],
    ("Field", "Type", "Req.", "Rules / default"): [2500, 1050, 750, 5060],
    ("Input", "Type", "Behavior"): [2700, 1800, 4860],
    ("Code", "Meaning", "Retryable"): [2400, 5280, 1680],
    ("Entity", "Identity", "Key fields"): [1800, 2700, 4860],
    ("Value", "Condition", "Client interpretation"): [1800, 4860, 2700],
    ("ID", "Requirement", "Priority"): [900, 7260, 1200],
    ("Component", "Responsibility", "Cloudflare-first implementation"): [1800, 3000, 4560],
    ("Data", "Authority / storage", "Default policy", "Stale behavior"): [1500, 2460, 2700, 2700],
    ("Constraint", "Published platform envelope", "V1 application rule"): [1900, 2900, 4560],
    ("Area", "Requirement"): [2100, 7260],
    ("Route", "Meaning", "Must not do"): [1800, 3660, 3900],
    ("Condition", "Initial threshold", "Response"): [2700, 2700, 3960],
    ("Phase", "Scope", "Exit criteria"): [1800, 3960, 3600],
    ("Question", "Decision for V1"): [3420, 5940],
    ("Binding / route", "Purpose", "Required for V1"): [2400, 3900, 3060],
}


@dataclass
class SourceDocument:
    metadata: dict[str, str]
    body_lines: list[str]


def parse_source(path: Path) -> SourceDocument:
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    if not lines or lines[0].strip() != "---":
        raise ValueError("Source must begin with YAML-like front matter")

    metadata: dict[str, str] = {}
    end = None
    for index, line in enumerate(lines[1:], start=1):
        if line.strip() == "---":
            end = index
            break
        key, separator, value = line.partition(":")
        if not separator:
            raise ValueError(f"Invalid front-matter line: {line!r}")
        metadata[key.strip()] = value.strip()
    if end is None:
        raise ValueError("Unclosed front matter")

    required = {
        "document_label",
        "title",
        "subtitle",
        "status",
        "owner",
        "target",
        "version",
        "last_updated",
        "decision_summary",
    }
    missing = required - metadata.keys()
    if missing:
        raise ValueError(f"Missing front-matter keys: {sorted(missing)}")
    return SourceDocument(metadata=metadata, body_lines=lines[end + 1 :])


def remove_body_content(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    fonts = run._element.rPr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_shading(element, fill: str) -> None:
    properties = element.get_or_add_pPr() if element.tag == qn("w:p") else element.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, *, side: str, color: str, size: int, space: int = 1) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    border = borders.find(qn(f"w:{side}"))
    if border is None:
        border = OxmlElement(f"w:{side}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|https?://[^\s]+)")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([run_properties, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str, *, bold: bool = False) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            run.bold = bold
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Courier New", 9.5)
            run.font.color.rgb = RGBColor.from_string(NAVY)
            run.bold = bold
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            suffix = ""
            while token and token[-1] in ".,;:)":
                suffix = token[-1] + suffix
                token = token[:-1]
            add_hyperlink(paragraph, token, token)
            if suffix:
                paragraph.add_run(suffix)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold = bold


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_text_size(paragraph, points: float) -> None:
    """Apply a size to regular and hyperlink runs in one paragraph."""
    half_points = str(round(points * 2))
    for run in paragraph._p.iter(qn("w:r")):
        run_properties = run.find(qn("w:rPr"))
        if run_properties is None:
            run_properties = OxmlElement("w:rPr")
            run.insert(0, run_properties)
        for tag in ("w:sz", "w:szCs"):
            size = run_properties.find(qn(tag))
            if size is None:
                size = OxmlElement(tag)
                run_properties.append(size)
            size.set(qn("w:val"), half_points)


def configure_table_width(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {TABLE_WIDTH_DXA}: {widths}")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.allow_autofit = False
    table_properties = table._tbl.tblPr

    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.insert(0, table_width)
    table_width.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")

    table_indent = table_properties.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:w"), "0")
    table_indent.set(qn("w:type"), "dxa")

    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{side}"))
        if border is None:
            border = OxmlElement(f"w:{side}")
            borders.append(border)
        border.set(qn("w:val"), "nil")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        cannot_split = row_properties.find(qn("w:cantSplit"))
        if cannot_split is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:w"), str(widths[index]))
            cell_width.set(qn("w:type"), "dxa")

            margins = cell_properties.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell_properties.append(margins)
            for side, value in (("top", 90), ("left", 120), ("bottom", 90), ("right", 120)):
                margin = margins.find(qn(f"w:{side}"))
                if margin is None:
                    margin = OxmlElement(f"w:{side}")
                    margins.append(margin)
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")


def choose_table_widths(headers: list[str]) -> list[int]:
    key = tuple(headers)
    if key in TABLE_WIDTHS_BY_HEADERS:
        return TABLE_WIDTHS_BY_HEADERS[key]

    count = len(headers)
    if count == 2:
        return [3000, 6360]
    if count == 3:
        return [2400, 3000, 3960]
    if count == 4:
        return [2100, 1800, 1800, 3660]
    base = TABLE_WIDTH_DXA // count
    widths = [base] * count
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    return widths


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = len(rows[0])
    if any(len(row) != column_count for row in rows):
        raise ValueError(f"Inconsistent Markdown table row width: {rows[:2]}")

    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = document.styles["Normal Table"]
    configure_table_width(table, choose_table_widths(rows[0]))

    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)

    for row_index, row_data in enumerate(rows):
        fill = PALE_BLUE if row_index == 0 else (ROW_ALT if row_index % 2 == 0 else WHITE)
        for cell, value in zip(table.rows[row_index].cells, row_data):
            set_shading(cell._tc, fill)
            paragraph = cell.paragraphs[0]
            clear_paragraph(paragraph)
            paragraph.style = document.styles["Normal"]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(paragraph, value, bold=row_index == 0)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.2


def next_numbering_id(numbering) -> int:
    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    return max(ids, default=0) + 1


def create_decimal_abstract_numbering(document: Document) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level_type = OxmlElement("w:multiLevelType")
    multi_level_type.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level_type)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    paragraph_properties.extend([tabs, indent])
    level.extend([start, number_format, level_text, justification, paragraph_properties])
    abstract.append(level)
    numbering.append(abstract)
    return abstract_id


def new_numbering_instance(document: Document, abstract_id: int) -> int:
    numbering = document.part.numbering_part.element
    number_id = next_numbering_id(numbering)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    # LibreOffice (used by the render QA pipeline) may continue numbering
    # across distinct w:num instances that share one abstract definition.
    # Make the restart explicit so every Markdown ordered-list block begins at 1.
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    number.append(level_override)
    numbering.append(number)
    return number_id


def apply_numbering(paragraph, number_id: int) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    number_properties = paragraph_properties.find(qn("w:numPr"))
    if number_properties is None:
        number_properties = OxmlElement("w:numPr")
        paragraph_properties.append(number_properties)
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(number_id))
    number_properties.extend([level, number])


def add_code_block(document: Document, lines: Iterable[str]) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    set_shading(paragraph._p, ROW_ALT)
    set_paragraph_border(paragraph, side="left", color=BLUE, size=14, space=4)

    line_list = list(lines)
    for index, line in enumerate(line_list):
        run = paragraph.add_run(line)
        set_run_font(run, "Courier New", 8.5)
        run.font.color.rgb = RGBColor.from_string(NAVY)
        if index < len(line_list) - 1:
            run.add_break(WD_BREAK.LINE)


def add_cover(document: Document, metadata: dict[str, str]) -> None:
    label = document.add_paragraph(style="Normal")
    label.paragraph_format.space_before = Pt(28)
    label.paragraph_format.space_after = Pt(18)
    run = label.add_run(metadata["document_label"])
    set_run_font(run, "Aptos", 14)
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(SLATE)

    title = document.add_paragraph(style="Title")
    add_inline(title, metadata["title"])
    set_paragraph_border(title, side="bottom", color=BLUE, size=12, space=4)

    subtitle = document.add_paragraph(style="Subtitle")
    add_inline(subtitle, metadata["subtitle"])

    metadata_rows = [
        ["Field", "Value"],
        ["Status", metadata["status"]],
        ["Owner", metadata["owner"]],
        ["Target", metadata["target"]],
        ["Version", metadata["version"]],
        ["Last updated", metadata["last_updated"]],
    ]
    add_table(document, metadata_rows)

    decision = document.add_paragraph(style="Normal")
    decision.paragraph_format.space_before = Pt(4)
    decision.paragraph_format.space_after = Pt(8)
    prefix = decision.add_run("Decision summary: ")
    prefix.bold = True
    add_inline(decision, metadata["decision_summary"])


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_markdown_body(document: Document, lines: list[str]) -> None:
    decimal_abstract_id = create_decimal_abstract_numbering(document)
    index = 0
    active_number_id: int | None = None
    compact_sources = False

    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            active_number_id = None
            index += 1
            continue

        if stripped.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            if index >= len(lines):
                raise ValueError("Unclosed code block")
            add_code_block(document, code_lines)
            active_number_id = None
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            active_number_id = None
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            paragraph = document.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, text)
            if level == 1 and text.startswith("Appendix"):
                paragraph.paragraph_format.page_break_before = True
            if level == 1:
                compact_sources = text.startswith("Appendix C")
            active_number_id = None
            index += 1
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:].strip())
            if compact_sources:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(1)
                set_paragraph_text_size(paragraph, 9.5)
            active_number_id = None
            index += 1
            continue

        numbered_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered_match:
            if active_number_id is None:
                active_number_id = new_numbering_instance(document, decimal_abstract_id)
            paragraph = document.add_paragraph(style="List Number")
            apply_numbering(paragraph, active_number_id)
            add_inline(paragraph, numbered_match.group(1))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if (
                not next_line
                or next_line.startswith("#")
                or next_line.startswith("|")
                or next_line.startswith("- ")
                or next_line.startswith("```")
                or re.match(r"^\d+\.\s+", next_line)
            ):
                break
            paragraph_lines.append(next_line)
            index += 1
        paragraph = document.add_paragraph(style="Normal")
        add_inline(paragraph, " ".join(paragraph_lines))
        active_number_id = None


def update_core_properties(
    document: Document, metadata: dict[str, str], source_sha256: str
) -> None:
    properties = document.core_properties
    properties.title = metadata["title"]
    properties.subject = "Product requirements and Cloudflare Workers hosting decision"
    properties.author = "Open Bikeshare MCP project"
    properties.keywords = "MCP, GBFS, bikeshare, Cloudflare Workers, Durable Objects"
    properties.comments = (
        "Generated from docs/Open_Bikeshare_MCP_PRD.md using the retained original "
        f"DOCX template. Source SHA-256: {source_sha256}"
    )
    properties.modified = datetime.now()


def update_template_chrome(document: Document, metadata: dict[str, str]) -> None:
    """Keep retained header/footer furniture while refreshing dated text."""
    for section in document.sections:
        for text_node in section.footer._element.iter(qn("w:t")):
            if text_node.text and re.search(
                r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4}\b",
                text_node.text,
            ):
                text_node.text = re.sub(
                    r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4}\b",
                    metadata["last_updated"],
                    text_node.text,
                )


def request_field_update(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def validate_document(document: Document) -> None:
    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.50
    assert round(section.page_height.inches, 2) == 11.00
    assert all(
        round(value.inches, 2) == 1.00
        for value in (
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        )
    )
    assert len(document.tables) >= 10
    assert any(paragraph.style.name == "Heading 1" for paragraph in document.paragraphs)


def build(source_path: Path, template_path: Path, output_path: Path) -> None:
    source = parse_source(source_path)
    source_sha256 = hashlib.sha256(source_path.read_bytes()).hexdigest()
    if template_path.resolve() == output_path.resolve():
        raise ValueError("Output path must differ from retained template")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    document = Document(output_path)
    remove_body_content(document)
    add_cover(document, source.metadata)
    add_markdown_body(document, source.body_lines)
    update_core_properties(document, source.metadata, source_sha256)
    update_template_chrome(document, source.metadata)
    request_field_update(document)
    validate_document(document)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.source.resolve(), args.template.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
